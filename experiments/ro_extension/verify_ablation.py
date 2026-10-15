"""Verify shared candidates, selection, paired replay and inventory identities."""
import json
import numpy as np
import pandas as pd
from constraint_ablation import HOME, pick, digest
from extend_with_ro import LABELS

def main():
    out=HOME/'outputs';design=json.loads((out/'frozen_design.json').read_text(encoding='utf-8'))
    candidates=pd.read_csv(out/'candidate_validation.csv')
    summary=pd.read_csv(out/'summary.csv')
    assert len(candidates)==design['library_size'] and len(summary)==12
    model=np.load(out/'shared_model.npz')
    assert design['refit'] is True and design['final_fit_periods']==list(range(1,61))
    assert model['train_y_all60'].shape==(60,6000)
    assert digest(model['train_y_all60'])==design['training_y_sha256']
    assert digest(model['final_labels'])==design['final_labels_sha256']
    assert digest(model['final_relative_residual'])==design['final_residual_sha256']
    import sys
    sys.path.insert(0,str(HOME.parent/'rolling_dro_v3'))
    from model_v2 import Forecaster
    refit=Forecaster(model['train_y_all60'],model['final_labels'],**design['forecast'])
    np.testing.assert_allclose(refit.slope,model['final_slope'],rtol=1e-12,atol=1e-10)
    np.testing.assert_allclose(refit.intercept,model['final_intercept'],rtol=1e-12,atol=1e-10)
    for key,field in [('relative_residual','residual_sha256'),('forecasts_validation','forecast_validation_sha256'),('labels','labels_sha256'),('initial','initial_sha256')]:
        assert digest(model[key])==design[field]
    common=None
    for selected in design['selected']:
        arm,method=selected['arm'],selected['method']
        chosen=pick(candidates[candidates.method==method],arm)
        assert int(chosen.candidate)==selected['candidate']
        f=pd.read_csv(out/f'detail_{arm}_{method}.csv.gz')
        assert len(f)==168000 and not f.duplicated(['period','item_id']).any()
        d,a,s,e,q=[f[c].to_numpy().reshape(28,6000) for c in ('demand','available','sales','ending','order')]
        np.testing.assert_array_equal(s,np.minimum(a,d));np.testing.assert_array_equal(e,a-s)
        np.testing.assert_array_equal(a[0],model['initial'])
        np.testing.assert_array_equal(a[1:],e[:-1]+q[:-1]);np.testing.assert_array_equal(q.ravel(),f.planned_order_h1)
        assert np.all(q>=0) and np.all(q==np.floor(q))
        for h in range(4):np.testing.assert_allclose(f[f'forecast_h{h+1}'].to_numpy().reshape(28,6000),model['forecasts_test'][:,h],rtol=1e-14)
        if common is None:common=d
        else:np.testing.assert_array_equal(d,common)
        r=summary[(summary.arm==arm)&(summary.method==method)].iloc[0]
        np.testing.assert_allclose([r.fill_rate,r.positive_stock_ratio,r.inventory_demand_ratio],[s.sum()/d.sum(),np.mean(e>0),e.sum()/d.sum()],rtol=1e-12)
    print('Ablation passed: common library/model/forecasts, training-only selection, all 12 full traces and summaries.')
    latest=HOME/'new_positive/outputs_v4'
    current=pd.read_csv(latest/'summary.csv').query("scope=='all_28_periods'").set_index('method')
    paired=summary[summary.arm=='positive'].set_index('method')
    np.testing.assert_allclose(current.loc[list(LABELS),'inventory_demand_ratio'],paired.loc[list(LABELS),'inventory_demand_ratio'],rtol=1e-12)
    for method in LABELS:
        a=pd.read_csv(latest/f'detail_{method}.csv.gz',usecols=['ending','order','sales'])
        b=pd.read_csv(out/f'detail_positive_{method}.csv.gz',usecols=['ending','order','sales'])
        np.testing.assert_array_equal(a,b)
        orders=pd.read_csv(latest/f'orders_{method}.csv').iloc[:,1:].to_numpy()
        np.testing.assert_array_equal(orders.T.ravel(),a.order)
    current_workbook=latest/'replenishment_comparison.xlsx'
    if current_workbook.exists():
        from openpyxl import load_workbook
        book=load_workbook(current_workbook,read_only=True,data_only=True)
        for i,method in enumerate(LABELS,5):
            r=current.loc[method]
            np.testing.assert_allclose([book['结论'].cell(i,4).value,book['结论'].cell(i,7).value,book['结论'].cell(i,9).value],
                [r.fill_rate,r.positive_stock_ratio,r.inventory_demand_ratio],rtol=1e-11)
        book.close()
    print('Full-60 final refit, six-policy exports, full order matrices and current workbook passed.')
    workbook=out/'新旧及单因素总对照.xlsx'
    if workbook.exists():
        from openpyxl import load_workbook
        book=load_workbook(workbook,read_only=True,data_only=True)
        for sheet,records in [('单因素对照',summary),('历史新旧两版',pd.concat([
            pd.read_csv(HOME.parent/v/'with_ro/outputs_v4/summary.csv').query("scope=='all_28_periods'")
            for v in ('demand_fill_v3','rolling_dro_v3')],ignore_index=True))]:
            sh=book[sheet]
            for i,r in enumerate(records.itertuples(),5):
                np.testing.assert_allclose([sh.cell(i,5).value,sh.cell(i,8).value,sh.cell(i,10).value],
                    [r.fill_rate,r.positive_stock_ratio,r.inventory_demand_ratio],rtol=1e-11)
        book.close();print('Combined workbook: 24 historical/controlled rows and formula caches passed.')

if __name__=='__main__':main()
