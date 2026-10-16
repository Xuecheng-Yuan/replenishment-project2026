"""Independent saved-output verification, with no raw test files or solver needed."""
import hashlib, json
import numpy as np
import pandas as pd
from extend_with_ro import ROOT, VERSIONS, LABELS, box

def verify(version, check_workbook=True):
    old=version==VERSIONS[0]
    base=ROOT/'experiments'/version/'outputs_v4'
    out=base.parent/'with_ro/outputs_v4'
    summary=pd.read_csv(out/'summary.csv').set_index(['method','scope'])
    assert len(summary)==12 and set(summary.index.get_level_values(0))==set(LABELS)
    original=pd.read_csv(base/'summary.csv').set_index(['method','scope'])
    for col in original.select_dtypes('number').columns:
        np.testing.assert_allclose(summary.loc[original.index,col],original[col],rtol=1e-12,atol=1e-7)
    params=json.loads((out/'frozen_parameters.json').read_text(encoding='utf-8'))
    chosen=params['selected']['RO']
    val=pd.read_csv(out/'ro_validation.csv')
    if old: eligible=val[(val.forward_fill>=params['forward_control_target'])&(val.reset_fill>=.95)]
    else: eligible=val[val.reset_service.between(.95,.955)]
    assert len(eligible)>0
    assert np.isclose(eligible.reset_inventory_ratio.min(),chosen['reset_inventory_ratio'])
    assert np.any(np.isclose(eligible.gamma,chosen['gamma']))
    for method in LABELS:
        f=pd.read_csv(out/f'detail_{method}.csv.gz')
        assert len(f)==168000 and not f.duplicated(['period','item_id']).any()
        d,a,s,e,q=[f[c].to_numpy().reshape(28,6000) for c in ('demand','available','sales','ending','order')]
        np.testing.assert_array_equal(s,np.minimum(a,d))
        np.testing.assert_array_equal(e,a-s)
        np.testing.assert_array_equal(a[1:],e[:-1]+q[:-1])
        np.testing.assert_array_equal(f.planned_order_h1,q.ravel())
        assert np.all(q>=0) and np.all(q==np.floor(q))
        result=summary.loc[(method,'all_28_periods')]
        np.testing.assert_allclose([result.fill_rate,result.positive_stock_ratio,result.inventory_demand_ratio],
            [s.sum()/d.sum(),np.mean(e>0),e.sum()/d.sum()],rtol=1e-12)
        if method!='RO':
            assert hashlib.sha256((out/f'detail_{method}.csv.gz').read_bytes()).digest()==hashlib.sha256((base/f'detail_{method}.csv.gz').read_bytes()).digest()
        else:
            for h in range(1,5):
                forecast=f[f'forecast_h{h}'].to_numpy()
                upper=forecast*chosen['upper_relative']
                expected=chosen['beta']*upper if old else np.floor(upper)+1
                np.testing.assert_allclose(f[f'target_h{h}'],expected,rtol=1e-12,atol=1e-8)
        del f
    full=pd.read_csv(out/'decision_comparison_full.csv.gz')
    assert len(full)==168000 and len([c for c in full if c.startswith('order_')])==6
    for method in LABELS:
        orders=pd.read_csv(out/f'orders_{method}.csv')
        np.testing.assert_array_equal(orders.iloc[:,1:].to_numpy().T.ravel(),full['order_'+method])
    if check_workbook:
        from openpyxl import load_workbook
        book=load_workbook(out/'replenishment_comparison.xlsx',read_only=True,data_only=True)
        sh=book['结论']
        for i,method in enumerate(LABELS,5):
            r=summary.loc[(method,'all_28_periods')]
            np.testing.assert_allclose([sh.cell(i,4).value,sh.cell(i,7).value,sh.cell(i,9).value],
                [r.fill_rate,r.positive_stock_ratio,r.inventory_demand_ratio],rtol=1e-11)
        latest=[row for row in book['最近一期全部SKU'].iter_rows(min_row=5,values_only=True) if row[0] is not None]
        sample=[row for row in book['完整周期样本'].iter_rows(min_row=5,values_only=True) if row[0] is not None]
        assert len(latest)==6000 and len(sample)==336
        reference=full[full.period==28]
        for i,method in enumerate(LABELS):
            np.testing.assert_array_equal([row[9+i] for row in latest],reference['order_'+method])
        book.close()
    manifest=out/'SHA256.json'
    if manifest.exists():
        for name,digest in json.loads(manifest.read_text(encoding='utf-8')).items():
            assert hashlib.sha256((out/name).read_bytes()).hexdigest()==digest,name
    print(version+': six methods, original traces, calibration, flow, RO endpoints, Excel and hashes passed')

if __name__=='__main__':
    for version in VERSIONS: verify(version)
