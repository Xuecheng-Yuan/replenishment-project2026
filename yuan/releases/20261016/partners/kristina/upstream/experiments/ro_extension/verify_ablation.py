"""Verify shared candidates, selection, paired replay and inventory identities."""
import json
import numpy as np
import pandas as pd
from constraint_ablation import HOME, FOLDS, pick, digest
from extend_with_ro import LABELS, box


def verify_rolling_calibration(out,design,saved,model):
    """Independently combine fold primitives and replay selected controllers."""
    folds=design['calibration_folds'];train=saved['train_y_all60'];init=saved['initial']
    assert len(folds)==3
    scored=[]
    for f,(cut,end) in zip(folds,FOLDS):
        assert f['fit_periods']==list(range(1,cut+1)) and f['validation_periods']==list(range(cut+1,end+1))
        assert digest(train[:cut])==f['fit_y_sha256'] and digest(train[cut:end])==f['validation_y_sha256']
        i=f['fold'];assert i in (1,2,3);scored+=f['validation_periods']
        for key,field in [('labels','labels_sha256'),('residual','residual_sha256'),('fc','forecast_sha256'),('warm_fc','warm_forecast_sha256')]:
            assert digest(saved[f'fold{i}_{key}'])==f[field]
        forecaster=model.Forecaster(train[:cut],saved[f'fold{i}_labels'],**design['forecast'])
        forecasts=model.make_forecasts(forecaster,train[cut:end],cut,history=train[:cut])
        np.testing.assert_allclose(forecasts,saved[f'fold{i}_fc'],rtol=1e-12,atol=1e-9)
        np.testing.assert_allclose(model.make_forecasts(forecaster,train[:cut],0),saved[f'fold{i}_warm_fc'],rtol=1e-12,atol=1e-9)
    assert scored==list(range(25,61))
    by_fold=pd.read_csv(out/'candidate_validation_by_fold.csv')
    aggregate=pd.read_csv(out/'candidate_validation.csv').set_index('candidate')
    assert len(by_fold)==3*design['library_size'] and not by_fold.duplicated(['candidate','fold']).any()
    grouped=by_fold.groupby('candidate')
    sums=grouped[['sales_units','demand_units','ending_inventory_units','positive_stock_sku_periods','total_sku_periods']].sum()
    np.testing.assert_allclose(aggregate.loc[sums.index,'fill_rate'],sums.sales_units/sums.demand_units,rtol=1e-12)
    np.testing.assert_allclose(aggregate.loc[sums.index,'positive_stock_ratio'],sums.positive_stock_sku_periods/sums.total_sku_periods,rtol=1e-12)
    np.testing.assert_allclose(aggregate.loc[sums.index,'inventory_demand_ratio'],sums.ending_inventory_units/sums.demand_units,rtol=1e-12)
    assert (grouped.fold.nunique()==3).all() and (sums.total_sku_periods==36*6000).all()
    forecast=pd.read_csv(out/'forecast_validation.csv');forecast_folds=pd.read_csv(out/'forecast_validation_by_fold.csv')
    assert len(forecast)==36 and len(forecast_folds)==108
    joined=forecast_folds.groupby(['blend','window','prior'])[['absolute_error','demand_units']].sum()
    for r in forecast.itertuples():
        totals=joined.loc[(r.blend,r.window,r.prior)]
        np.testing.assert_allclose(r.wape,totals.absolute_error/totals.demand_units,rtol=1e-12)
        for f,(cut,end) in zip(folds,FOLDS):
            fm=model.Forecaster(train[:cut],saved[f"fold{f['fold']}_labels"],blend=r.blend,window=int(r.window),prior=r.prior)
            predicted=model.rolling_predictions(fm,train[cut:end],cut,history=train[:cut])[:,0]
            actual=forecast_folds[(forecast_folds.blend==r.blend)&(forecast_folds.window==r.window)&(forecast_folds.prior==r.prior)&(forecast_folds.fold==f['fold'])].iloc[0]
            np.testing.assert_allclose([np.abs(predicted-train[cut:end]).sum(),train[cut:end].sum()],[actual.absolute_error,actual.demand_units],rtol=1e-12)
    chosen=forecast.sort_values('wape',kind='stable').iloc[0]
    assert all(chosen[k]==design['forecast'][k] for k in ('blend','window','prior'))
    from run_v2 import ss_quantiles
    hot=pd.read_csv(out/'sku_summary.csv').query("arm=='positive' and method=='sS'").set_index('item_id').category.eq('爆品').to_numpy(dtype=int)
    for s in design['selected']:
        recipe=design['recipe'][s['candidate']];method=s['method']
        for f,(cut,end) in zip(folds,FOLDS):
            i=f['fold'];labels=saved[f'fold{i}_labels'];residual=saved[f'fold{i}_residual']
            high=None
            if method=='RO':low=np.full(6000,box(residual,recipe['gamma'])[1])
            elif method=='sS':
                low=ss_quantiles(residual,labels,recipe['p_s'])*recipe['scale']
                high=ss_quantiles(residual,labels,recipe['p_S'])*recipe['scale']
            else:
                group=hot if method=='DRO_category' else np.zeros(6000,int) if method=='DRO_global' else labels
                low=model.risk_levels(residual,group,recipe['beta'],recipe['radius'])[0]*recipe['scale']
            warm=model.simulate(train[:cut],init,saved[f'fold{i}_warm_fc'],low,high,positive_guard=True,keep=True)
            state=warm['ending'][-1]+warm['orders'][-1]
            result=model.simulate(train[cut:end],state,saved[f'fold{i}_fc'],low,high,positive_guard=True)['summary']
            record=by_fold[(by_fold.candidate==s['candidate'])&(by_fold.fold==i)].iloc[0]
            np.testing.assert_allclose(state.sum(),record.validation_initial_available_units,rtol=1e-12)
            for key in ('sales_units','demand_units','ending_inventory_units','positive_stock_sku_periods'):
                np.testing.assert_allclose(result[key],record[key],rtol=1e-12)
    cache=HOME.parent/'rolling_dro_v3/.cache/period_data.npz'
    if cache.exists():
        source=np.load(cache)
        np.testing.assert_array_equal(train,source['train_y'])
        for f,(cut,end) in zip(folds,FOLDS):assert digest(source['train_x'][:cut].astype(float))==f['fit_x_sha256']
    print('Three rolling folds passed: ranges, causal forecasts, 108 forecast scores, all candidate aggregates and 36 selected-fold replays.')

def main():
    out=HOME/'outputs';design=json.loads((out/'frozen_design.json').read_text(encoding='utf-8'))
    candidates=pd.read_csv(out/'candidate_validation.csv')
    summary=pd.read_csv(out/'summary.csv')
    assert len(candidates)==design['library_size'] and len(summary)==12
    model=np.load(out/'shared_model.npz')
    assert design['refit'] is True and design['final_fit_periods']==list(range(1,61))
    assert model['train_y_all60'].shape==(60,6000)
    assert digest(model['train_y_all60'])==design['training_y_sha256']
    assert digest(model['final_labels'])==design['final_labels_sha256']
    assert digest(model['final_relative_residual'])==design['final_residual_sha256']
    import sys
    sys.path.insert(0,str(HOME.parent/'rolling_dro_v3'))
    import model_v2 as engine
    refit=engine.Forecaster(model['train_y_all60'],model['final_labels'],**design['forecast'])
    np.testing.assert_allclose(refit.slope,model['final_slope'],rtol=1e-12,atol=1e-10)
    np.testing.assert_allclose(refit.intercept,model['final_intercept'],rtol=1e-12,atol=1e-10)
    assert digest(model['initial'])==design['initial_sha256']
    verify_rolling_calibration(out,design,model,engine)
    common=None
    for selected in design['selected']:
        arm,method=selected['arm'],selected['method']
        chosen=pick(candidates[candidates.method==method],arm)
        assert int(chosen.candidate)==selected['candidate']
        f=pd.read_csv(out/f'detail_{arm}_{method}.csv.gz')
        assert len(f)==168000 and not f.duplicated(['period','item_id']).any()
        d,a,s,e,q=[f[c].to_numpy().reshape(28,6000) for c in ('demand','available','sales','ending','order')]
        np.testing.assert_array_equal(s,np.minimum(a,d));np.testing.assert_array_equal(e,a-s)
        np.testing.assert_array_equal(a[0],model['initial'])
        np.testing.assert_array_equal(a[1:],e[:-1]+q[:-1]);np.testing.assert_array_equal(q.ravel(),f.planned_order_h1)
        assert np.all(q>=0) and np.all(q==np.floor(q))
        for h in range(4):np.testing.assert_allclose(f[f'forecast_h{h+1}'].to_numpy().reshape(28,6000),model['forecasts_test'][:,h],rtol=1e-14)
        if common is None:common=d
        else:np.testing.assert_array_equal(d,common)
        r=summary[(summary.arm==arm)&(summary.method==method)].iloc[0]
        np.testing.assert_allclose([r.fill_rate,r.positive_stock_ratio,r.inventory_demand_ratio],[s.sum()/d.sum(),np.mean(e>0),e.sum()/d.sum()],rtol=1e-12)
    print('Ablation passed: common library/model/forecasts, training-only selection, all 12 full traces and summaries.')
    np.testing.assert_allclose(engine.make_forecasts(refit,common,0),model['forecasts_test'],rtol=1e-12,atol=1e-9)
    original=pd.read_csv(HOME.parent/'rolling_dro_v3/outputs_v4/detail_sS.csv.gz',usecols=['demand'])
    np.testing.assert_array_equal(original.demand.to_numpy().reshape(28,6000),common)
    latest=HOME/'new_positive/outputs_v4'
    current=pd.read_csv(latest/'summary.csv').query("scope=='all_28_periods'").set_index('method')
    paired=summary[summary.arm=='positive'].set_index('method')
    np.testing.assert_allclose(current.loc[list(LABELS),'inventory_demand_ratio'],paired.loc[list(LABELS),'inventory_demand_ratio'],rtol=1e-12)
    for method in LABELS:
        a=pd.read_csv(latest/f'detail_{method}.csv.gz',usecols=['ending','order','sales'])
        b=pd.read_csv(out/f'detail_positive_{method}.csv.gz',usecols=['ending','order','sales'])
        np.testing.assert_array_equal(a,b)
        orders=pd.read_csv(latest/f'orders_{method}.csv').iloc[:,1:].to_numpy()
        np.testing.assert_array_equal(orders.T.ravel(),a.order)
    current_workbook=latest/'replenishment_comparison.xlsx'
    if current_workbook.exists():
        from openpyxl import load_workbook
        book=load_workbook(current_workbook,read_only=True,data_only=True)
        for i,method in enumerate(LABELS,5):
            r=current.loc[method]
            np.testing.assert_allclose([book['结论'].cell(i,4).value,book['结论'].cell(i,7).value,book['结论'].cell(i,9).value],
                [r.fill_rate,r.positive_stock_ratio,r.inventory_demand_ratio],rtol=1e-11)
        segments=pd.read_csv(latest/'summary_by_segment.csv')
        for i,r in enumerate(segments.itertuples(),5):
            np.testing.assert_allclose([book['爆品长尾结果'].cell(i,j).value for j in (6,9,11)],
                [r.fill_rate,r.positive_stock_ratio,r.inventory_demand_ratio],rtol=1e-11)
        folds=pd.read_csv(out/'candidate_validation_by_fold.csv')
        chosen={r['method']:r['candidate'] for r in design['selected'] if r['arm']=='positive'}
        selected=folds[folds.apply(lambda r:r.candidate==chosen[r.method],axis=1)]
        for i,r in enumerate(selected.itertuples(),5):
            np.testing.assert_allclose([book['三轮验证'].cell(i,j).value for j in (5,6,7)],
                [r.fill_rate,r.positive_stock_ratio,r.inventory_demand_ratio],rtol=1e-11)
        book.close()
    print('Full-60 final refit, six-policy exports, full order matrices and current workbook passed.')
    workbook=out/'新旧及单因素总对照.xlsx'
    if workbook.exists():
        from openpyxl import load_workbook
        book=load_workbook(workbook,read_only=True,data_only=True)
        historical=pd.concat([
            pd.read_csv(HOME.parent/v/'with_ro/outputs_v4/summary.csv').query("scope=='all_28_periods'")
            for v in ('demand_fill_v3','rolling_dro_v3')]+[
                pd.read_csv(HOME/'history/single_holdout_32_28/summary.csv').query("arm=='positive'")],ignore_index=True)
        for sheet,records in [('单因素对照',summary),('历史新旧两版',historical)]:
            sh=book[sheet]
            for i,r in enumerate(records.itertuples(),5):
                np.testing.assert_allclose([sh.cell(i,5).value,sh.cell(i,8).value,sh.cell(i,10).value],
                    [r.fill_rate,r.positive_stock_ratio,r.inventory_demand_ratio],rtol=1e-11)
        selected_folds=pd.read_csv(out/'selected_validation_by_fold.csv')
        for i,r in enumerate(selected_folds.itertuples(),5):
            np.testing.assert_allclose([book['逐轮验证'].cell(i,j).value for j in (8,11,13)],
                [r.fill_rate,r.positive_stock_ratio,r.inventory_demand_ratio],rtol=1e-11)
        forecast=pd.read_csv(out/'forecast_validation.csv')
        for i,r in enumerate(forecast.itertuples(),5):
            np.testing.assert_allclose(book['预测参数验证'].cell(i,6).value,r.wape,rtol=1e-11)
        book.close();print('Combined workbook: 30 historical/controlled rows and formula caches passed.')

if __name__=='__main__':main()
