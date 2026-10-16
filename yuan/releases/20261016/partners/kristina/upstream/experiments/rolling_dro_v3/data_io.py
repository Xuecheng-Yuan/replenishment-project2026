from __future__ import annotations
import re
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from config import Config
TRAIN_FILE="replenishment_train.xlsx"; TEST_FILE="replenishment_test.xlsx"; CACHE_FILE="period_data.npz"
FEATURE_RE=re.compile(r"^x\d+$",re.IGNORECASE); PERIOD_RE=re.compile(r"^Period_(\d+)$",re.IGNORECASE)
def _period_number(name):
    match=PERIOD_RE.match(name)
    if not match: raise ValueError(f"Unexpected worksheet name: {name}")
    return int(match.group(1))
def _read_workbook(path):
    print(f"[data] streaming {path.name} (first read can take several minutes)",flush=True)
    book=load_workbook(path,read_only=True,data_only=True)
    names=sorted(book.sheetnames,key=_period_number)
    if not names: raise ValueError(f"No Period_* worksheets found in {path}")
    if [_period_number(name) for name in names]!=list(range(1,len(names)+1)):
        raise ValueError(f"Period numbering has gaps in {path}")
    item_id=initial_inventory=category=is_hot=None
    y_values=[]; x_values=[]
    for name in names:
        sheet=book[name]
        rows=sheet.iter_rows(values_only=True)
        header=list(next(rows))
        index={str(value):j for j,value in enumerate(header)}
        feature_idx=[index[f'x{k}'] for k in range(1,21)]
        demand_cols=[j for j,c in enumerate(header) if str(c).startswith('demand_')]
        if len(demand_cols)!=1: raise ValueError(f"Expected one demand column in {name}")
        raw=list(rows)
        ids=np.fromiter((r[index['item_id']] for r in raw),dtype=np.int64,count=len(raw))
        initial=np.fromiter((r[index['initial_inventory']] for r in raw),dtype=np.float32,count=len(raw))
        cats=np.asarray([str(r[index['category']]) for r in raw])
        demand=np.fromiter((r[demand_cols[0]] for r in raw),dtype=np.float32,count=len(raw))
        features=np.asarray([[r[j] for j in feature_idx] for r in raw],dtype=np.float32)
        if item_id is None:
            item_id,initial_inventory,category=ids,initial,cats
            is_hot=category=='爆品'
        elif not (np.array_equal(ids,item_id) and np.array_equal(initial,initial_inventory)
                  and np.array_equal(cats,category)):
            raise ValueError(f"item_id, category, or initial_inventory differs in {name}")
        if not (np.isfinite(demand).all() and np.isfinite(features).all() and (demand>=0).all()):
            raise ValueError(f"Invalid demand or feature in {name}")
        y_values.append(demand); x_values.append(features)
        if len(y_values)%10==0: print(f'[data] {path.name}: {len(y_values)}/{len(names)} periods',flush=True)
    book.close()
    return {"y":np.stack(y_values),"x":np.stack(x_values),"item_id":item_id,"initial_inventory":initial_inventory,"is_hot":is_hot}
def load_period_data(config:Config,refresh_cache:bool=False):
    config.cache_dir.mkdir(parents=True,exist_ok=True); cache_path=config.cache_dir/CACHE_FILE
    if cache_path.exists() and not refresh_cache:
        data=np.load(cache_path,allow_pickle=False); print(f"[data] cache hit: {cache_path}"); return {key:data[key] for key in data.files}
    train_path=config.data_dir/TRAIN_FILE; test_path=config.data_dir/TEST_FILE
    if not train_path.exists() or not test_path.exists(): raise FileNotFoundError(f"Expected {TRAIN_FILE} and {TEST_FILE} under {config.data_dir}")
    train=_read_workbook(train_path); test=_read_workbook(test_path)
    if not np.array_equal(train["item_id"],test["item_id"]): raise ValueError("train and test item_id orders differ")
    if not np.array_equal(train["is_hot"],test["is_hot"]): raise ValueError("train and test category assignments differ")
    if not np.allclose(train["initial_inventory"],test["initial_inventory"]): raise ValueError("train and test initial_inventory differ")
    data={"train_y":train["y"],"train_x":train["x"],"test_y":test["y"],"test_x":test["x"],"item_id":train["item_id"],"initial_inventory":train["initial_inventory"],"is_hot":train["is_hot"]}
    np.savez_compressed(cache_path,**data); print(f"[data] cached compact arrays: {cache_path}"); return data
