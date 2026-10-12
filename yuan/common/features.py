"""common.features: actual model/processing implementation, not a legacy runner alias.

Refactored mechanically from the existing project. See matching docs for
information timing, historical results, variants and limitations.
"""
from __future__ import annotations
import sys
import tempfile
from pathlib import Path
HERE=Path(__file__).resolve().parent
ROOT=Path(__file__).resolve().parents[1]
if str(HERE) in sys.path:sys.path.remove(str(HERE))
if str(ROOT) not in sys.path:sys.path.insert(0,str(ROOT))
PROJECT_ROOT=ROOT
DATA_DIR=ROOT/'data'
OUTPUT_DIR=ROOT/'results/generated'
import re
from dataclasses import dataclass
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
X_COLUMNS = [f'x{i}' for i in range(1, 21)]
DEMAND_LAGS = (1, 2, 3, 6, 12)
ROLLING_WINDOWS = (3, 6, 12)

@dataclass(frozen=True)
class PanelData:
    item_ids: np.ndarray
    categories: np.ndarray
    initial_inventory: np.ndarray
    demand: np.ndarray
    x: np.ndarray

@dataclass(frozen=True)
class SupervisedData:
    features: np.ndarray
    target: np.ndarray | None
    feature_names: list[str]
    target_periods: np.ndarray

def load_panel(file_path: str | Path) -> PanelData:
    """Read all workbook periods into compact NumPy arrays."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f'找不到数据文件: {path}')
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: list[tuple[int, str]] = []
        for sheet_name in workbook.sheetnames:
            match = re.fullmatch('Period_(\\d+)', sheet_name, flags=re.IGNORECASE)
            if match:
                sheets.append((int(match.group(1)), sheet_name))
        sheets.sort()
        if [period for period, _ in sheets] != list(range(1, len(sheets) + 1)):
            raise ValueError('工作簿周期表必须从 Period_1 连续编号。')
        item_ids: np.ndarray | None = None
        categories: np.ndarray | None = None
        initial_inventory: np.ndarray | None = None
        demand_by_period: list[np.ndarray] = []
        x_by_period: list[np.ndarray] = []
        for period, sheet_name in sheets:
            rows = workbook[sheet_name].iter_rows(values_only=True)
            header = [str(value).strip() if value is not None else '' for value in next(rows)]
            positions = {name: index for index, name in enumerate(header)}
            demand_name = f'demand_{period}'
            required = {'item_id', 'category', 'initial_inventory', demand_name, *X_COLUMNS}
            missing = required - set(positions)
            if missing:
                raise ValueError(f'{sheet_name} 缺少字段: {sorted(missing)}')
            records = list(rows)
            current_ids = np.asarray([row[positions['item_id']] for row in records])
            current_demand = np.asarray([row[positions[demand_name]] for row in records], dtype=float)
            current_x = np.asarray([[row[positions[column]] for column in X_COLUMNS] for row in records], dtype=float)
            if item_ids is None:
                item_ids = current_ids
                categories = np.asarray([row[positions['category']] for row in records])
                initial_inventory = np.asarray([row[positions['initial_inventory']] for row in records], dtype=float)
            elif not np.array_equal(item_ids, current_ids):
                raise ValueError(f'{sheet_name} 的 item_id 顺序或集合不一致。')
            demand_by_period.append(current_demand)
            x_by_period.append(current_x)
    finally:
        workbook.close()
    if item_ids is None or categories is None or initial_inventory is None:
        raise ValueError('工作簿没有可用周期。')
    demand = np.column_stack(demand_by_period)
    x = np.stack(x_by_period, axis=1)
    if not np.isfinite(demand).all() or not np.isfinite(x).all():
        raise ValueError('需求或x特征包含缺失/无穷值。')
    if (demand < 0).any() or (initial_inventory < 0).any():
        raise ValueError('需求和初始库存不能为负数。')
    return PanelData(item_ids, categories, initial_inventory, demand, x)

def build_supervised_data(demand: np.ndarray, x: np.ndarray, categories: np.ndarray, target_periods: np.ndarray, include_target: bool=True) -> SupervisedData:
    """Build period-major rows using information available before each target.

    For target period ``t``, every demand and anonymous x feature comes from
    period ``t-1`` or earlier. ``t`` may equal the observed history length to
    construct the next, not-yet-observed forecast row.
    """
    demand = np.asarray(demand, dtype=float)
    x = np.asarray(x, dtype=float)
    target_periods = np.asarray(target_periods, dtype=int)
    if demand.ndim != 2 or x.ndim != 3 or x.shape[:2] != demand.shape:
        raise ValueError('demand 必须是二维，x 必须是对应的三维数组。')
    if len(categories) != demand.shape[0]:
        raise ValueError('categories 行数与SKU数不一致。')
    if target_periods.size == 0 or target_periods.min() < max(DEMAND_LAGS):
        raise ValueError('目标周期必须至少有12期历史。')
    maximum = demand.shape[1] - 1 if include_target else demand.shape[1]
    if target_periods.max() > maximum:
        raise ValueError('目标周期超出可构造范围。')
    category_values = sorted({str(value) for value in categories})
    category_map = {value: index for index, value in enumerate(category_values)}
    category_code = np.asarray([category_map[str(value)] for value in categories], dtype=float)
    rows: list[np.ndarray] = []
    targets: list[np.ndarray] = []
    for period in target_periods:
        lag_values = [demand[:, period - lag] for lag in DEMAND_LAGS]
        rolling_means = [demand[:, period - window:period].mean(axis=1) for window in ROLLING_WINDOWS]
        rolling_stds = [demand[:, period - window:period].std(axis=1, ddof=1) for window in ROLLING_WINDOWS]
        lag_1, lag_2, lag_3 = (lag_values[0], lag_values[1], lag_values[2])
        mean_3, mean_6 = (rolling_means[0], rolling_means[1])
        x_lag_1 = x[:, period - 1, :]
        x_lag_2 = x[:, period - 2, :]
        block = np.column_stack([*lag_values, *rolling_means, *rolling_stds, lag_1 - lag_2, lag_1 - lag_3, (lag_1 + 1.0) / (mean_6 + 1.0), (lag_1 - lag_3) / 2.0, np.log1p(lag_1), np.log1p(mean_3), x_lag_1, x_lag_1 - x_lag_2, category_code, np.full(demand.shape[0], period + 1, dtype=float)])
        rows.append(block)
        if include_target:
            targets.append(demand[:, period])
    names = [f'demand_lag_{lag}' for lag in DEMAND_LAGS]
    names += [f'demand_mean_{window}' for window in ROLLING_WINDOWS]
    names += [f'demand_std_{window}' for window in ROLLING_WINDOWS]
    names += ['demand_diff_1', 'demand_diff_3', 'lag1_to_mean6', 'recent_slope']
    names += ['log_demand_lag_1', 'log_demand_mean_3']
    names += [f'{column}_lag_1' for column in X_COLUMNS]
    names += [f'{column}_change_1' for column in X_COLUMNS]
    names += ['category_code', 'period_index']
    return SupervisedData(features=np.vstack(rows).astype(np.float32), target=np.concatenate(targets).astype(np.float32) if include_target else None, feature_names=names, target_periods=target_periods)
